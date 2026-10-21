"""DOCX-экспорт отчёта из Markdown (D-03)."""

from __future__ import annotations

import io

from docx import Document


def markdown_to_docx(markdown: str) -> bytes:
    """Собрать .docx из Markdown-отчёта (заголовки, списки, абзацы)."""
    document = Document()
    in_code_block = False
    for raw_line in markdown.splitlines():
        line = raw_line.rstrip()
        if line.strip() == "```":
            in_code_block = not in_code_block
            continue
        if in_code_block:
            document.add_paragraph(line)
        elif line.startswith("### "):
            document.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            document.add_heading(line[2:], level=1)
        elif line.startswith("- "):
            document.add_paragraph(line[2:], style="List Bullet")
        elif line.strip():
            document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
